"""Workbook schema and missing-data regression checks."""
from pathlib import Path
import shutil
import tempfile
import unittest

import numpy as np
from openpyxl import load_workbook

from fatigue.imp.experimental_data import import_experimental_data_grouped_by_material
from fatigue.imp.schema import material_id, dataset_column
from fatigue.eval.calibration_store import input_fingerprint

ROOT = Path(__file__).resolve().parents[1]


class DatasetImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.path = self.folder / '8_AISI316L.xlsx'
        shutil.copy(ROOT / 'Experimental_Data/AISI316L.xlsx', self.path)

    def test_filename_prefix_does_not_change_calibration_identity(self):
        baseline = import_experimental_data_grouped_by_material(ROOT / 'Experimental_Data')
        current = import_experimental_data_grouped_by_material(self.folder)
        self.assertEqual(input_fingerprint(baseline, 'AISI316L'), input_fingerprint(current, 'AISI316L'))
        self.assertEqual(material_id('2_CX18'), 'CX18')
        self.assertEqual(material_id('Ti6Al4V_andrea_(AXIAL)'), 'Ti6Al4V_andrea_(AXIAL)')

    def test_header_mapping_is_explicit_and_position_independent(self):
        self.assertEqual(dataset_column(['Mean AXIAL load (N)', 'Mean AXIAL stress  (MPa)'], 'mean_ax'),
                         'Mean AXIAL stress  (MPa)')
        with self.assertRaises(KeyError):
            dataset_column(['Mean AXIAL load (N)'], 'mean_ax')

    def test_missing_elastic_property_preserves_stress_and_runouts(self):
        w = load_workbook(self.path)
        s = w['Summary']
        for row in s:
            if str(row[0].value).lower() == 'average':
                row[5].value = 'Null'
        w.save(self.path)
        block = import_experimental_data_grouped_by_material(self.folder)['AISI316L']
        self.assertIsNone(block['material']['nu'])
        self.assertEqual(len(block['cases']), 44)
        self.assertTrue(any(c['meta']['test_status'] == 'runout' for c in block['cases'].values()))
        case = next(iter(block['cases'].values()))
        self.assertTrue(np.isfinite(case[1][1]['stress']).all())
        self.assertTrue(np.isnan(case[1][1]['strain']).all())

    def test_axial_phase_is_the_reference(self):
        w = load_workbook(self.path)
        s = w['Dataset']
        c = s.max_column + 1
        s.cell(1, c, 'Phase AXIAL load (°)')
        for row in range(2, s.max_row + 1):
            s.cell(row, c, 30)
        w.save(self.path)
        block = import_experimental_data_grouped_by_material(self.folder)['AISI316L']
        case = next(iter(block['cases'].values()))
        self.assertEqual(case['meta']['phase_rad_deg'], -30)
        self.assertEqual(case['meta']['phase_tor_deg'], -30)
        self.assertEqual(block['source']['format'], 'revised')
