"""Overlay layer: as-supplied workbooks plus declared corrections reproduce the revised dataset.

The five workbooks edited after supply are kept byte-identical to their source; their corrections
are shipped with the package. Each overlay is found by the SHA-256 of the file it was written against,
so it follows a copy of that file and never applies to an edited one.
"""

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from multifat.imp.experimental_data import (
    import_experimental_data_grouped_by_material,
    load_material_workbook,
)
from multifat.imp.overlays import OverlayError, apply_overlay_to_workbook, overlay_path
from multifat.report.subsets import dataset_fingerprint

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "Experimental_Data"
OVERLAID = ["2_CX18", "3_25CrMo4", "4_30NCD16", "7_982FA", "8_AISI316L"]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.mark.parametrize("stem", OVERLAID)
def test_overlay_names_the_as_supplied_bytes(stem):
    path = overlay_path(FOLDER / f"{stem}.xlsx")
    assert path is not None
    overlay = json.loads(path.read_text(encoding="utf-8"))
    assert overlay["base_sha256"] == _sha256(FOLDER / f"{stem}.xlsx")
    assert overlay["operations"], "an overlay with no operations should not exist"


@pytest.mark.parametrize("stem", OVERLAID)
def test_overlay_reproduces_expected_fingerprint(stem):
    overlay = json.loads(overlay_path(FOLDER / f"{stem}.xlsx").read_text(encoding="utf-8"))
    data = load_material_workbook(FOLDER / f"{stem}.xlsx")
    (block,) = data.values()
    assert block["overlay"]["sha256"] is not None
    assert block["overlay"]["operations"] == len(overlay["operations"])
    assert dataset_fingerprint(block) == overlay["expected_fingerprint"]


def test_aisi316l_counts_match_provenance():
    block = load_material_workbook(FOLDER / "8_AISI316L.xlsx")["AISI316L"]
    assert len(block["cases"]) == 39
    assert len(block["runouts"]) == 13


def test_overlay_off_gives_the_as_supplied_dataset():
    on = load_material_workbook(FOLDER / "8_AISI316L.xlsx")["AISI316L"]
    off = load_material_workbook(FOLDER / "8_AISI316L.xlsx", overlay=False)["AISI316L"]
    assert off["overlay"] is None
    assert len(off["cases"]) == 31
    assert on["material"]["E_MPa"] == 193000.0
    assert off["material"]["E_MPa"] != on["material"]["E_MPa"]


def test_no_overlay_for_workbooks_without_one():
    block = load_material_workbook(FOLDER / "1_Ti6Al4V.xlsx")["Ti6Al4V"]
    assert block["overlay"] is None


def test_folder_import_applies_overlays_by_default():
    data = import_experimental_data_grouped_by_material(FOLDER)
    assert data["CX18"]["material"]["nu"] == 0.3
    assert data["AISI316L"]["overlay"]["name"] == "8_AISI316L.overlay.json"


def test_tampered_from_value_aborts(tmp_path):
    wb = load_workbook(FOLDER / "8_AISI316L.xlsx", data_only=True)
    wb["Summary"]["E2"].value = 1.0
    with pytest.raises(OverlayError, match="expected"):
        apply_overlay_to_workbook(wb, FOLDER / "8_AISI316L.xlsx")


def test_renamed_copy_keeps_its_overlay(tmp_path):
    copy = tmp_path / "some_other_name.xlsx"
    copy.write_bytes((FOLDER / "8_AISI316L.xlsx").read_bytes())
    (block,) = load_material_workbook(copy).values()
    assert block["overlay"]["name"] == "8_AISI316L.overlay.json"
    assert len(block["cases"]) == 39


def test_modified_workbook_is_not_overlaid(tmp_path):
    wb = load_workbook(FOLDER / "8_AISI316L.xlsx", data_only=True)
    copy = tmp_path / "8_AISI316L.xlsx"
    wb.save(copy)  # same content, different bytes: no overlay is written against it
    assert overlay_path(copy) is None
    assert apply_overlay_to_workbook(wb, copy) is None
