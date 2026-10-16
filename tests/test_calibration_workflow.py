"""Regression checks for JSON history and explicitly censored observations."""
import copy
import contextlib
import io
import json
from pathlib import Path
import runpy
import shutil
import tempfile
import unittest

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from fatigue.imp import experimental_data as exp
from fatigue.imp.observations import test_status as status_from_flag, failures_only, observation_counts
from fatigue.eval.calibration import export_calibration
from fatigue.eval.calibration_store import (curve_settings, input_fingerprint, load_calibration,
                                           save_record, validate_record)
from fatigue.eval.fat_cycles import fatigue_plots
from fatigue.report.export_csv import save_fatigue_csv

ROOT = Path(__file__).resolve().parents[1]


class CalibrationWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = exp.import_experimental_data_grouped_by_material(ROOT / "Experimental_Data")
        cls.material = "AISI316L"
        cls.settings = curve_settings(3, 6)
        cls.script0 = runpy.run_path(str(ROOT / "0_calibrate_material_params.py"))
        cls.compare = runpy.run_path(str(ROOT / "2_compare_results.py"))

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def export(self, data=None, settings=None):
        return export_calibration(data=data or self.data, material=self.material, models=["BP"],
                                  parameters={}, settings=settings or self.settings, folder=self.folder)

    def load(self, **kwargs):
        return load_calibration(data=kwargs.pop("data", self.data), material=self.material,
                                model="BP", settings=self.settings, folder=self.folder, **kwargs)

    def test_round_trip_precision_units_metadata_and_history(self):
        first, p1 = self.export()
        repeated, repeated_path = self.export()
        self.assertEqual(repeated_path, p1)
        self.assertEqual(repeated, first)
        second, p2 = self.export(settings=curve_settings(3, 6, 1.0))
        self.assertNotEqual(p1, p2)
        self.assertTrue(p1.is_file())
        record, curve, path = load_calibration(data=self.data, material=self.material,
            model="BP", settings=curve_settings(3, 6, 1.0), folder=self.folder)
        self.assertEqual(path, p2)
        self.assertEqual(record["material_id"], self.material)
        self.assertEqual(record["curve_units"]["N"], "cycle")
        self.assertEqual(record["counts"]["runout"], 13)
        np.testing.assert_array_equal(curve["series"]["N"], second["curves"]["BP"]["series"]["N"])
        self.assertEqual(curve["fit"], second["curves"]["BP"]["fit"])
        self.assertEqual(self.load(calibration_id=first["calibration_id"])[2], p1)
        # Selection follows record time, not filesystem modification time.
        p1.touch()
        self.assertEqual(load_calibration(data=self.data, material=self.material,
            model="BP", settings=curve_settings(3, 6, 1.0), folder=self.folder)[2], p2)
        with self.assertRaises(FileExistsError):
            save_record(first, self.folder)

    def test_equivalence_ignores_censored_counts_rows_and_runtime_provenance(self):
        record, path = self.export()
        repeated = copy.deepcopy(record)
        repeated["calibration_id"] += "_repeat"
        repeated["software"]["git_commit"] = "another-commit"
        repeated["software"]["tracked_worktree_dirty"] = not repeated["software"]["tracked_worktree_dirty"]
        repeated["source"]["workbook_sha256"] = "0" * 64
        repeated["counts"]["runout"] += 1
        repeated["design_observations"][next(iter(repeated["design_observations"]))]["excel_row"] = 999
        self.assertEqual(save_record(repeated, self.folder), path)
        self.assertEqual(len(list((self.folder / self.material).glob("*.json"))), 1)
        self.assertEqual(json.loads(path.read_text()), record)

    def test_changed_calibration_output_creates_history_even_with_same_inputs(self):
        record, path = self.export()
        changed = copy.deepcopy(record)
        changed["calibration_id"] += "_changed"
        changed["curves"]["BP"]["fit"]["A"] *= 1.01
        self.assertNotEqual(save_record(changed, self.folder), path)

    def test_missing_and_incompatible_calibrations(self):
        with self.assertRaisesRegex(ValueError, "Run python 0_"):
            self.load()
        record, _ = self.export()
        changed = copy.deepcopy(self.data)
        changed[self.material]["material"]["Yield_strength_MPa"] += 1
        with self.assertRaisesRegex(ValueError, "incompatible"):
            self.load(data=changed, calibration_id=record["calibration_id"])
        changed = copy.deepcopy(self.data)
        next(c for c in changed[self.material]["cases"].values()
             if c["meta"]["test_status"] == "failure")["meta"]["Nf_cycles"] += 1
        with self.assertRaisesRegex(ValueError, "No compatible"):
            self.load(data=changed)
        with self.assertRaisesRegex(ValueError, "Invalid material/calibration identifier"):
            self.load(calibration_id="../invalid")

    def test_invalid_records_never_become_successful_history(self):
        record, path = self.export()
        bad = copy.deepcopy(record)
        bad["curves"]["BP"]["fit"]["A_surv"] = float("inf")
        with self.assertRaisesRegex(ValueError, "Invalid calibration"):
            save_record(bad, self.folder)
        path.write_text('{"schema_version":1,"curves":NaN}')
        with self.assertRaisesRegex(ValueError, "Cannot read calibration"):
            self.load()
        self.assertFalse(list(self.folder.rglob("*.tmp")))

    def test_missing_required_coefficient_is_rejected(self):
        record, _ = self.export()
        record["curves"]["FS"] = record["curves"]["BP"]
        with self.assertRaisesRegex(ValueError, "requires parameter k_FS"):
            validate_record(record)

    def test_formulation_and_settings_changes_reject_stored_curves(self):
        record, path = self.export()
        with self.assertRaisesRegex(ValueError, "No compatible"):
            load_calibration(data=self.data, material=self.material, model="BP", settings=curve_settings(4, 8), folder=self.folder)
        record["formulation_fingerprints"]["BP"] = "0" * 64
        path.write_text(json.dumps(record))
        with self.assertRaisesRegex(ValueError, "incompatible"):
            self.load(calibration_id=record["calibration_id"])

    def test_failure_only_curve_and_fingerprint_ignore_runouts(self):
        failures = failures_only(self.data)
        a, _ = self.export(failures)
        b, _ = export_calibration(data=self.data, material=self.material, models=["BP"],
            parameters={}, settings=self.settings, folder=self.folder / "with-runouts")
        self.assertEqual(a["curves"], b["curves"])
        self.assertEqual(a["input_fingerprint"], b["input_fingerprint"])
        self.assertEqual(a["counts"]["runout"], 0)
        self.assertEqual(b["counts"]["runout"], 13)
        reordered = copy.deepcopy(self.data)
        block = reordered[self.material]
        block["cases"] = {1000+i: c for i, c in enumerate(reversed(list(block["cases"].values())))}
        self.assertEqual(input_fingerprint(reordered, self.material), a["input_fingerprint"])
        self.assertIsNotNone(self.load(data=failures))

    def test_all_five_optimisations_ignore_runouts_and_match_previous_code(self):
        planes = np.array([np.eye(3), [[0, 0, 1], [0, 1, 0], [-1, 0, 0]]])
        def fit(data):
            return self.script0["calibrate_material_params"](data=data, material_name=self.material,
                R_list=planes, n_grid_1d=5, n_refine_1d=1)
        a, b = fit(failures_only(self.data)), fit(self.data)
        baseline = json.loads((ROOT / "tests/fixtures/parameter_calibration_baseline.json").read_text())
        for method, param in [("FS", "k_FS"), ("FI", "k_FI"), ("LI", "L_LI"),
                              ("MGSE_ZHU", "k_MGSE"), ("MSWT", "a_MSWT")]:
            self.assertEqual(a[method], b[method])
            self.assertAlmostEqual(a[method][param], baseline[method][param], places=12)
            for key in ["A", "b", "A_surv", "STD"]:
                np.testing.assert_allclose(a[method]["fit"][key], baseline[method]["fit"][key], rtol=1e-12, atol=1e-12)

    def test_explicit_status_import_unknown_and_invalid_observations(self):
        self.assertEqual(status_from_flag("Yes"), "runout")
        self.assertEqual(status_from_flag("no"), "failure")
        self.assertEqual(status_from_flag(None), "unknown")
        source = ROOT / "Experimental_Data/AISI316L.xlsx"
        target = self.folder / source.name
        shutil.copy2(source, target)
        # This temporary fixture uses cached Summary values when appending rows.
        wb = load_workbook(target, data_only=True)
        ws = wb["Dataset"]
        row = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
        columns = {str(c.value): c.column - 1 for c in ws[1]}
        row[columns["Run out (Yes/No)"]] = "Yes"
        ws.append(row)
        row[columns["Run out (Yes/No)"]] = "?"
        ws.append(row)
        row[columns["Number of cycles to failure"]] = 0
        ws.append(row)
        wb.save(target)
        wb.close()
        imported = exp.import_experimental_data_grouped_by_material(self.folder)
        counts = observation_counts(imported[self.material])
        self.assertEqual(counts, {"failure": 31, "runout": 14, "unknown": 1, "invalid_excluded": 1})
        self.assertEqual(input_fingerprint(imported, self.material), input_fingerprint(self.data, self.material))

    def test_runouts_have_no_reconstructed_errors_or_statistics(self):
        path = save_fatigue_csv(out_dir=str(self.folder), method_name="BP", material_name=self.material,
            N_exp=[1e5, 2e6], CP=[1., .5], N_pred=[2e5, 3e6], DP_fit=[1.1, .6],
            test_status=["failure", "runout"])
        result = self.compare["read_results_csv"](Path(path))
        self.assertTrue(result.loc[1, ["Error_ln", "Error_ln_dp", "Error_log10_dp", "DP_fit"]].isna().all())
        stats = self.compare["compute_stats_table"](result)
        self.assertTrue(stats.n.eq(1).all())

    def test_runout_markers_arrows_and_legends(self):
        f1, ax1, f2, ax2 = fatigue_plots([1., .5], [1e5, 2e6], [2e5, 3e6], 10, -.2,
            test_status=["failure", "runout"], show=False)
        try:
            for ax in [ax1, ax2]:
                self.assertTrue(any("lower bound" in label for label in ax.get_legend_handles_labels()[1]))
                self.assertTrue(any(getattr(text, "arrow_patch", None) for text in ax.texts))
        finally:
            plt.close(f1)
            plt.close(f2)

    def test_comparison_png_and_tikz_preserve_censoring(self):
        data = pd.DataFrame({"Method": ["BP", "BP"], "Nf_exp": [1e5, 2e6],
                            "Nf_expected": [2e5, 3e6], "Test_status": ["failure", "runout"]})
        png, tex = self.folder / "plot.png", self.folder / "plot.tex"
        self.compare["plot_expected_vs_experimental"](data, png, tex, self.folder / "data",
                                                      self.material, {}, .2, [3, 5], False)
        self.assertTrue(png.is_file())
        self.assertIn("Run-out (experimental life lower bound)", tex.read_text())
        self.assertIn(r"\draw[->", tex.read_text())
        self.assertEqual(np.loadtxt(self.folder / "data/scatter/BP.dat").shape, (2,))

    def test_script0_to_script1_json_integration(self):
        with contextlib.redirect_stdout(io.StringIO()):
            paths = self.script0["main"](["--material", self.material, "--models", "FS",
                "--data-folder", str(ROOT / "Experimental_Data"), "--calibrations-dir", str(self.folder),
                "--n1", "3", "--n2", "6", "--n-grid", "3", "--n-refine", "1", "--plane-stride", "1"])
            record = json.loads(paths[0].read_text())
            self.assertIn("k_FS", record["parameters"])
            self.assertEqual(record["parameter_units"]["k_FS"], "1")
            runner = runpy.run_path(str(ROOT / "1_run_material.py"))
            runner["main"].__globals__.update(DATA_FOLDER=str(ROOT / "Experimental_Data"),
                CALIBRATIONS_DIR=str(self.folder), MATERIAL_NAME=self.material, MODEL_NAME="FS",
                DCMESH_N1=3, DCMESH_N2=6, RESULTS_DIR=str(self.folder), MAKE_PLOTS=False)
            runner["main"]()
        result = pd.read_csv(self.folder / "FS_AISI316L.csv")
        self.assertTrue(result.Calibration_id.eq(record["calibration_id"]).all())
        self.assertEqual(result.Test_status.value_counts().to_dict(), {"failure": 31, "runout": 13})


if __name__ == "__main__":
    unittest.main()
