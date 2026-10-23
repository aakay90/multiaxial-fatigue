"""Compare workbook structure and imported inputs without editing measured data.

Ported from upstream tools/audit_dataset_migration.py (A. Chiocca, 904c619) to the multifat
importer: failures and run-outs are separate blocks there, the fingerprint is
multifat.report.subsets.dataset_fingerprint, and Summary source rows are not exported.
"""

import argparse
import hashlib
import json
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from multifat.imp.experimental_data import (  # noqa: E402
    import_experimental_data_grouped_by_material,
)
from multifat.imp.schema import material_id, workbook_files  # noqa: E402
from multifat.report.subsets import dataset_fingerprint, json_value  # noqa: E402


def headers(path):
    w = load_workbook(path, read_only=True, data_only=True)
    value = {
        name: [str(c.value) for c in next(w[name].iter_rows()) if c.value is not None]
        for name in ("Summary", "Dataset")
    }
    w.close()
    return value


def counts(block):
    return {
        "failure": len(block["cases"]),
        "runout": len(block.get("runouts", {})),
        "excluded": len(block.get("excluded_observations", [])),
    }


def _key(case):
    m = case["meta"]
    return (m.get("specimen_id"), m["Nf_cycles"], m.get("test_status"))


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--old-dir", default="Experimental_Data/archive/pre-dataset-migration")
    parser.add_argument("--new-dir", default="Experimental_Data")
    parser.add_argument("--output", default="Results/current/Comparison/dataset-migration.json")
    args = parser.parse_args()
    old = import_experimental_data_grouped_by_material(args.old_dir)
    new = import_experimental_data_grouped_by_material(args.new_dir)
    old_paths = {material_id(p.stem): p for p in workbook_files(args.old_dir)}
    new_paths = {material_id(p.stem): p for p in workbook_files(args.new_dir)}
    records = []
    for material in sorted(old.keys() | new.keys()):
        before, after = old.get(material), new.get(material)
        record = {
            "material": material,
            "old_workbook": old_paths[material].name if before else None,
            "new_workbook": new_paths[material].name if after else None,
            "status": "retained" if before and after else "added" if after else "archived",
        }
        for label, block in [("old", before), ("new", after)]:
            if block:
                record[label + "_observations"] = counts(block)
                record[label + "_properties"] = block["material"]
        if after:
            record["identity_issues"] = after.get("identity_issues", [])
            record["property_issues"] = after.get("property_issues", {})
            record["new_workbook_sha256"] = hashlib.sha256(
                new_paths[material].read_bytes()
            ).hexdigest()
        if before and after:
            record["calibration_inputs_unchanged"] = dataset_fingerprint(
                before
            ) == dataset_fingerprint(after)
            first, second = headers(old_paths[material]), headers(new_paths[material])
            record["header_changes"] = {
                sheet: {
                    "removed": sorted(set(first[sheet]) - set(second[sheet])),
                    "added": sorted(set(second[sheet]) - set(first[sheet])),
                }
                for sheet in first
            }
            old_cases = {_key(c): c for c in before["cases"].values()}
            record["matching_stress_observations"] = sum(
                _key(c) in old_cases
                and all(
                    np.array_equal(c[1][s]["stress"], old_cases[_key(c)][1][s]["stress"])
                    for s in (1, 2)
                )
                for c in after["cases"].values()
            )
        records.append(record)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(json_value(records), indent=2, allow_nan=False) + "\n")
    print(f"Compared {len(old)} previous and {len(new)} current material IDs; wrote {output}")


if __name__ == "__main__":
    main()
